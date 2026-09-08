import io
import re
from typing import List, Optional
import openpyxl
from docx import Document
import pdfplumber

from app.core.schemas import BaiHoc

def parse_ppct_docx(file_bytes: bytes, default_subject: str = "Toán") -> List[BaiHoc]:
    """
    Parse Phân phối chương trình / Giáo án từ file Word (.docx).
    Tối ưu hóa khả năng nhận diện cột tên bài dạy cho mọi mẫu bảng của giáo viên.
    """
    doc = Document(io.BytesIO(file_bytes))
    items: List[BaiHoc] = []
    
    current_mon = default_subject

    # Quét lấy tên môn từ văn bản ngoài bảng
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "MÔN:" in txt.upper() or "MÔN HỌC:" in txt.upper() or "KẾ HOẠCH DẠY HỌC MÔN" in txt.upper():
            parts = re.split(r'MÔN\s*HỌC:|MÔN:|MÔN', txt, flags=re.IGNORECASE)
            if len(parts) > 1 and parts[1].strip():
                # Lấy tên môn (bỏ các ký tự rác)
                mon_clean = re.sub(r'[^a-zA-Zà-ỹÀ-Ỹ\s]', '', parts[1]).strip()
                if len(mon_clean) > 1:
                    current_mon = mon_clean

    for table in doc.tables:
        if not table.rows:
            continue

        col_tiet_idx = -1
        col_mon_idx = -1
        col_ten_idx = -1
        col_tb_idx = -1
        
        # 1. Quét tìm cột trong các hàng đầu tiên (Header detection)
        for r_idx, row in enumerate(table.rows[:5]):
            cells_text = [c.text.strip() for c in row.cells]
            for c_idx, text in enumerate(cells_text):
                t_lower = text.lower()
                if any(k in t_lower for k in ["tiết", "stt", "tuần", "t.số tiết"]):
                    if col_tiet_idx == -1: col_tiet_idx = c_idx
                elif any(k in t_lower for k in ["môn", "phân môn"]):
                    if col_mon_idx == -1: col_mon_idx = c_idx
                elif any(k in t_lower for k in ["tên bài", "nội dung", "bài dạy", "chủ đề", "bài học", "bài"]):
                    if col_ten_idx == -1: col_ten_idx = c_idx
                elif any(k in t_lower for k in ["thiết bị", "đồ dùng", "thtb", "đddh"]):
                    if col_tb_idx == -1: col_tb_idx = c_idx

        # 2. Nếu không tìm thấy cột Tên bài qua header, dùng thuật toán đoán cột theo độ dài chuỗi
        if col_ten_idx == -1 and len(table.rows) > 1:
            # Chọn cột có độ dài chuỗi văn bản trung bình dài nhất
            col_lengths = {}
            for row in table.rows[1:]:
                for c_idx, cell in enumerate(row.cells):
                    txt_len = len(cell.text.strip())
                    col_lengths[c_idx] = col_lengths.get(c_idx, 0) + txt_len
            if col_lengths:
                # Cột có tổng độ dài dài nhất thường là Tên Bài Dạy
                col_ten_idx = max(col_lengths, key=col_lengths.get)

        # 3. Đọc dữ liệu các hàng trong bảng
        for r_idx, row in enumerate(table.rows):
            cells_text = [c.text.strip() for c in row.cells]
            
            if col_ten_idx != -1 and col_ten_idx < len(cells_text):
                ten_bai = cells_text[col_ten_idx]
                
                # Bỏ qua các hàng chứa tiêu đề hoặc hàng rỗng
                if not ten_bai or any(k in ten_bai.lower() for k in ["tên bài", "nội dung bài", "tên bài dạy", "chủ đề / bài"]):
                    continue
                if ten_bai.isdigit() and len(ten_bai) < 4:
                    continue # Bỏ qua nếu dòng này chứa số tiết

                # Tiết PPCT
                tiet_val = len(items) + 1
                if col_tiet_idx != -1 and col_tiet_idx < len(cells_text):
                    t_str = cells_text[col_tiet_idx]
                    nums = re.findall(r'\d+', t_str)
                    if nums:
                        tiet_val = int(nums[0])

                # Môn học
                mon_val = current_mon
                if col_mon_idx != -1 and col_mon_idx < len(cells_text) and cells_text[col_mon_idx]:
                    m_txt = cells_text[col_mon_idx]
                    if not m_txt.isdigit() and len(m_txt) > 1:
                        mon_val = m_txt

                # Thiết bị
                thiet_bi_val = ""
                if col_tb_idx != -1 and col_tb_idx < len(cells_text):
                    thiet_bi_val = cells_text[col_tb_idx]

                items.append(BaiHoc(
                    mon=mon_val,
                    tiet_ppct=tiet_val,
                    ten_bai=ten_bai,
                    thiet_bi=thiet_bi_val
                ))

    return items

def parse_ppct_excel(file_bytes: bytes, default_subject: str = "Toán") -> List[BaiHoc]:
    """
    Parse Phân phối chương trình từ file Excel (.xlsx).
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    items: List[BaiHoc] = []

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return items

    col_tiet_idx = -1
    col_mon_idx = -1
    col_ten_idx = -1
    col_tb_idx = -1

    # 1. Đọc Header
    for r_idx, row in enumerate(rows[:5]):
        row_str = [str(c).strip() if c is not None else "" for c in row]
        for c_idx, text in enumerate(row_str):
            t_lower = text.lower()
            if any(k in t_lower for k in ["tiết", "stt", "tuần"]):
                if col_tiet_idx == -1: col_tiet_idx = c_idx
            elif any(k in t_lower for k in ["môn", "phân môn"]):
                if col_mon_idx == -1: col_mon_idx = c_idx
            elif any(k in t_lower for k in ["tên bài", "nội dung", "bài dạy", "chủ đề", "bài học", "bài"]):
                if col_ten_idx == -1: col_ten_idx = c_idx
            elif any(k in t_lower for k in ["thiết bị", "đồ dùng"]):
                if col_tb_idx == -1: col_tb_idx = c_idx

    # Fallback đoán cột Tên bài nếu không tìm thấy header
    if col_ten_idx == -1 and len(rows) > 1:
        col_lengths = {}
        for row in rows[1:]:
            for c_idx, cell in enumerate(row):
                if cell is not None:
                    txt_len = len(str(cell).strip())
                    col_lengths[c_idx] = col_lengths.get(c_idx, 0) + txt_len
        if col_lengths:
            col_ten_idx = max(col_lengths, key=col_lengths.get)

    # 2. Đọc dữ liệu
    for row in rows:
        row_str = [str(c).strip() if c is not None else "" for c in row]
        if col_ten_idx != -1 and col_ten_idx < len(row_str):
            ten_bai = row_str[col_ten_idx]
            if not ten_bai or any(k in ten_bai.lower() for k in ["tên bài", "nội dung bài", "bài dạy", "chủ đề"]):
                continue

            tiet_val = len(items) + 1
            if col_tiet_idx != -1 and col_tiet_idx < len(row_str):
                nums = re.findall(r'\d+', row_str[col_tiet_idx])
                if nums:
                    tiet_val = int(nums[0])

            mon_val = default_subject
            if col_mon_idx != -1 and col_mon_idx < len(row_str) and row_str[col_mon_idx]:
                mon_val = row_str[col_mon_idx]

            thiet_bi_val = row_str[col_tb_idx] if col_tb_idx != -1 and col_tb_idx < len(row_str) else ""

            items.append(BaiHoc(
                mon=mon_val,
                tiet_ppct=tiet_val,
                ten_bai=ten_bai,
                thiet_bi=thiet_bi_val
            ))

    return items

def parse_ppct_pdf(file_bytes: bytes, default_subject: str = "Toán") -> List[BaiHoc]:
    """
    Parse Phân phối chương trình từ file PDF (.pdf).
    """
    items: List[BaiHoc] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                if not table:
                    continue
                col_tiet_idx = -1
                col_mon_idx = -1
                col_ten_idx = -1
                col_tb_idx = -1

                for r_idx, row in enumerate(table):
                    row_str = [str(c).strip() if c is not None else "" for c in row]
                    if col_ten_idx == -1:
                        for c_idx, text in enumerate(row_str):
                            t_lower = text.lower()
                            if any(k in t_lower for k in ["tiết", "stt", "tuần"]):
                                col_tiet_idx = c_idx
                            elif any(k in t_lower for k in ["môn", "phân môn"]):
                                col_mon_idx = c_idx
                            elif any(k in t_lower for k in ["tên bài", "nội dung", "bài dạy", "chủ đề", "bài"]):
                                col_ten_idx = c_idx
                            elif any(k in t_lower for k in ["thiết bị", "đồ dùng"]):
                                col_tb_idx = c_idx
                        continue

                    if col_ten_idx != -1 and col_ten_idx < len(row_str):
                        ten_bai = row_str[col_ten_idx]
                        if not ten_bai:
                            continue

                        tiet_val = len(items) + 1
                        if col_tiet_idx != -1 and col_tiet_idx < len(row_str):
                            nums = re.findall(r'\d+', row_str[col_tiet_idx])
                            if nums:
                                tiet_val = int(nums[0])

                        mon_val = default_subject
                        if col_mon_idx != -1 and col_mon_idx < len(row_str) and row_str[col_mon_idx]:
                            mon_val = row_str[col_mon_idx]

                        thiet_bi_val = row_str[col_tb_idx] if col_tb_idx != -1 and col_tb_idx < len(row_str) else ""

                        items.append(BaiHoc(
                            mon=mon_val,
                            tiet_ppct=tiet_val,
                            ten_bai=ten_bai,
                            thiet_bi=thiet_bi_val
                        ))

    return items

def auto_parse_ppct(file_bytes: bytes, filename: str) -> List[BaiHoc]:
    """
    Tự động thử các parser (Word, Excel, PDF) để đọc file PPCT bất kể đuôi file bị đặt sai.
    """
    fn = filename.lower()
    parsers = []

    if fn.endswith(".xlsx") or fn.endswith(".xls"):
        parsers = [parse_ppct_excel, parse_ppct_docx, parse_ppct_pdf]
    elif fn.endswith(".pdf"):
        parsers = [parse_ppct_pdf, parse_ppct_docx, parse_ppct_excel]
    else:
        parsers = [parse_ppct_docx, parse_ppct_excel, parse_ppct_pdf]

    for parser in parsers:
        try:
            res = parser(file_bytes)
            if res:
                return res
        except Exception:
            continue

    return []
