import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def generate_docxtpl_template(template_path: str = "app/templates/mau_lich.docx"):
    os.makedirs(os.path.dirname(template_path), exist_ok=True)
    doc = Document()

    # Cấu hình lề
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Tiêu đề
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("LỊCH BÁO DẠY - TUẦN {{ tuan }}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph()

    # Bảng
    headers = ["Thứ", "Buổi", "Tiết TKB", "Môn học", "Tiết PPCT", "Tên bài dạy / Nội dung", "Thiết bị", "Ghi chú"]
    table = doc.add_table(rows=2, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Headers
    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
        for p in hdr_cells[idx].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.bold = True
                r.font.size = Pt(11)

        # Style header background
        tcPr = hdr_cells[idx]._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'EBF2F7')
        tcPr.append(shd)

    # Row 1 with docxtpl tags
    row1_cells = table.rows[1].cells
    row1_cells[0].text = "{% for r in danh_sach %}{{ r.thu }}"
    row1_cells[1].text = "{{ r.buoi }}"
    row1_cells[2].text = "{{ r.tiet_tkb }}"
    row1_cells[3].text = "{{ r.mon }}"
    row1_cells[4].text = "{{ r.tiet_ppct }}"
    row1_cells[5].text = "{{ r.ten_bai }}"
    row1_cells[6].text = "{{ r.thiet_bi }}"
    row1_cells[7].text = "{{ r.ghi_chu }}{% endfor %}"

    for idx, cell in enumerate(row1_cells):
        for p in cell.paragraphs:
            if idx in [0, 1, 2, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

    doc.save(template_path)
    print(f"✅ Đã tạo file template Word tại: {template_path}")

if __name__ == "__main__":
    generate_docxtpl_template()
